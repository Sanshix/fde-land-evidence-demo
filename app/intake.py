"""Read-only multiformat evidence intake. Never execute macros, formulas, links or Office apps."""

import hashlib
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pypdfium2
from PIL import Image

VERSION = "intake-v1"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
X = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def issue(code, location, message, severity="review"):
    return {
        "code": code,
        "location": location,
        "message": message,
        "severity": severity,
        "category": "source_document",
    }


def unit(digest, index, name, source, kind="native", **extra):
    return {
        "id": f"{digest[:16]}-{index}",
        "name": name,
        "source": str(source),
        "source_sha256": digest,
        "kind": kind,
        "path": None,
        "rotation": 0,
        "native_lines": [],
        "issues": [],
        **extra,
    }


def safe_xml(raw):
    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise ValueError("XML实体/DTD不受支持，不执行")
    return ET.fromstring(raw)


def open_package(raw):
    z = zipfile.ZipFile(io.BytesIO(raw))
    entries = z.infolist()
    if len(entries) > 5000 or sum(i.file_size for i in entries) > 120 * 1024 * 1024:
        z.close()
        raise ValueError("Office包内容过大")
    for i in entries:
        if i.flag_bits & 1 or Path(i.filename).is_absolute() or ".." in Path(i.filename).parts:
            z.close()
            raise ValueError("Office包加密或路径无效")
    return z


def image_unit(raw, digest, index, name, source, objects, location=None):
    im = Image.open(io.BytesIO(raw))
    if im.width * im.height > 50000000:
        raise ValueError("图片像素超过5000万")
    if getattr(im, "n_frames", 1) > 1:
        raise ValueError("多帧图片尚不支持，请转多页PDF；不静默取首帧")
    im.load()
    target = objects / (hashlib.sha256(raw).hexdigest() + ".png")
    if not target.exists():
        im.convert("RGB").save(target)
    return unit(digest, index, name, source, "image", path=str(target), location=location or name)


def embedded_images(z, digest, source, objects, prefix):
    units = []
    problems = []
    for name in z.namelist():
        if name.startswith(prefix + "/media/") and not name.endswith("/"):
            try:
                units.append(
                    image_unit(
                        z.read(name),
                        digest,
                        "embedded-" + str(len(units) + 1),
                        name,
                        source,
                        objects,
                        location=name,
                    )
                )
            except Exception as e:
                problems.append(
                    issue(
                        "unsupported_embedded_image", name, type(e).__name__ + ": " + str(e)[:160], "blocking"
                    )
                )
        elif "/embeddings/" in name or name.lower().endswith("vbaproject.bin"):
            problems.append(
                issue(
                    "embedded_object_not_executed",
                    name,
                    "嵌入对象/宏仅存证，未执行；需补充可读格式",
                    "blocking",
                )
            )
        elif name.endswith(".rels"):
            for link in safe_xml(z.read(name)):
                if link.get("TargetMode") == "External":
                    problems.append(
                        issue(
                            "external_relationship_not_fetched",
                            name,
                            "外部关系已记录，未读取目标：" + link.get("Target", "")[:250],
                        )
                    )
    return units, problems


def docx(raw, digest, source, objects):
    units = []
    problems = []
    with open_package(raw) as z:
        parts = [
            p
            for p in z.namelist()
            if p == "word/document.xml"
            or re.match(r"word/(header\d*|footer\d*|footnotes|endnotes|comments)\.xml$", p)
        ]
        if "word/document.xml" not in parts:
            raise ValueError("DOCX缺少正文")
        for part in parts:
            root = safe_xml(z.read(part))
            u = unit(digest, len(units) + 1, part, source, part=part, location=part)
            # Paragraph locators include table/cell identity. Office pagination is not fabricated.
            paragraph_index = 0

            def add(p, location, output=u):
                nonlocal paragraph_index
                paragraph_index += 1
                text = "".join(t.text or "" for t in p.iter("{" + W + "}t"))
                deleted = "".join(t.text or "" for t in p.iter("{" + W + "}delText"))
                hidden = any(n.tag == "{" + W + "}vanish" for n in p.iter())
                tracked = any(n.tag in {"{" + W + "}ins", "{" + W + "}del"} for n in p.iter())
                if text or deleted:
                    output["native_lines"].append(
                        {
                            "text": text,
                            "raw_text": text,
                            "locator": location + f"/p{paragraph_index}",
                            "deleted_text": deleted,
                            "hidden": hidden,
                            "tracked_change": tracked,
                        }
                    )
                if deleted or tracked:
                    problems.append(issue("tracked_change", location, "存在修订文字；未自动接受/拒绝修订"))
                if hidden:
                    problems.append(issue("hidden_text", location, "存在隐藏文字，已列入证据"))

            body = root.find("{" + W + "}body")
            parent = body if body is not None else root
            table_index = 0
            for child in parent:
                if child.tag == "{" + W + "}tbl":
                    table_index += 1
                    for ri, row in enumerate(child.findall("{" + W + "}tr"), 1):
                        for ci, cell in enumerate(row.findall("{" + W + "}tc"), 1):
                            loc = f"{part}/table{table_index}/r{ri}c{ci}"
                            for p in cell.iter("{" + W + "}p"):
                                add(p, loc)
                            if (
                                cell.find(".//{" + W + "}gridSpan") is not None
                                or cell.find(".//{" + W + "}vMerge") is not None
                            ):
                                problems.append(
                                    issue("merged_table_cell", loc, "合并单元格，保留物理单元格位置")
                                )
                else:
                    for p in [child] if child.tag == "{" + W + "}p" else child.iter("{" + W + "}p"):
                        add(p, part)
            units.append(u)
        extra, found = embedded_images(z, digest, source, objects, "word")
        units.extend(extra)
        problems.extend(found)
        problems.append(
            issue(
                "docx_pagination_not_rendered",
                "word/document.xml",
                "DOCX以段落/表格单元格定位，未渲染Word分页；不伪造页码",
            )
        )
        if any("/charts/" in p or "/diagrams/" in p for p in z.namelist()):
            problems.append(
                issue("drawing_not_rendered", "word", "图表/SmartArt未渲染，需原件补验", "blocking")
            )
    return units, problems


def xlsx(raw, digest, source, objects):
    units = []
    problems = []
    with open_package(raw) as z:
        # Read raw XML lexical values separately from openpyxl's typed interpretation.
        wbxml = safe_xml(z.read("xl/workbook.xml"))
        sheet_nodes = wbxml.findall(".//{" + X + "}sheet")
        relroot = safe_xml(z.read("xl/_rels/workbook.xml.rels"))
        rels = {r.get("Id"): r.get("Target") for r in relroot}
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=False, keep_links=False)
        try:
            for sheetnode in sheet_nodes:
                name = sheetnode.get("name")
                ws = wb[name]
                relationship = sheetnode.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                )
                target = rels[relationship]
                part = target.lstrip("/") if target.startswith("/") else "xl/" + target
                root = safe_xml(z.read(part))
                cells = root.findall(".//{" + X + "}sheetData/{" + X + "}row/{" + X + "}c")
                if len(cells) > 20000:
                    raise ValueError("单工作表非空单元格超过20000，需拆分")
                u = unit(
                    digest,
                    len(units) + 1,
                    f"工作表 · {name}",
                    source,
                    sheet=name,
                    part=part,
                    location=name,
                    hidden=ws.sheet_state != "visible",
                    merged_ranges=[str(r) for r in ws.merged_cells.ranges],
                )
                if u["hidden"]:
                    problems.append(issue("hidden_sheet", name, "隐藏工作表已纳入证据，未忽略"))
                for cell in cells:
                    coord = cell.get("r")
                    typed = ws[coord]
                    kind = cell.get("t", "n")
                    rawnode = cell.find("{" + X + "}v")
                    lexical = rawnode.text if rawnode is not None else None
                    formula = cell.find("{" + X + "}f")
                    formula_text = formula.text if formula is not None else None
                    text = (
                        "".join(t.text or "" for t in cell.iter("{" + X + "}t"))
                        if kind == "inlineStr"
                        else str(typed.value or "")
                        if kind == "s"
                        else lexical or ""
                    )
                    if formula is not None:
                        text = "=" + str(formula_text or "")
                    hidden = bool(
                        ws.row_dimensions[typed.row].hidden
                        or ws.column_dimensions[typed.column_letter].hidden
                    )
                    display = text
                    display_status = "literal"
                    fmt = typed.number_format
                    if formula is None and kind == "n" and lexical is not None:
                        if re.fullmatch(r"0+", fmt) and re.fullmatch(r"\d+", lexical):
                            display = lexical.zfill(len(fmt))
                            display_status = "zero_mask"
                        elif isinstance(typed.value, datetime):
                            display = typed.value.isoformat()
                            display_status = "date_interpretation_check_epoch"
                        elif fmt != "General":
                            display_status = "format_not_fully_rendered"
                    entry = {
                        "text": text,
                        "raw_text": text,
                        "locator": f"{name}!{coord}",
                        "cell": coord,
                        "raw_xml_value": lexical,
                        "cell_type": kind,
                        "number_format": fmt,
                        "display_value": display,
                        "display_status": display_status,
                        "formula": formula_text,
                        "cached_value": lexical if formula is not None else None,
                        "hidden": hidden,
                        "date_epoch": str(wb.epoch),
                    }
                    u["native_lines"].append(entry)
                    if hidden:
                        problems.append(issue("hidden_cell", entry["locator"], "隐藏行/列内容已纳入证据"))
                    if formula is not None:
                        problems.append(
                            issue(
                                "formula_not_executed",
                                entry["locator"],
                                "保留公式与缓存值；缓存可能陈旧，不能当已复算结果",
                                "blocking",
                            )
                        )
                    if kind == "n" and lexical and len(re.sub(r"[^0-9]", "", lexical)) >= 15:
                        problems.append(
                            issue(
                                "numeric_identifier_precision",
                                entry["locator"],
                                "长数字以原始XML词法值保存；若Excel已损失精度无法逆推出原编号",
                                "blocking",
                            )
                        )
                    if display_status == "format_not_fully_rendered":
                        problems.append(
                            issue(
                                "display_format_unresolved",
                                entry["locator"],
                                "自定义格式未完整渲染；原值与格式均保留",
                            )
                        )
                for merged in u["merged_ranges"]:
                    problems.append(
                        issue(
                            "merged_cells", name + "!" + merged, "保留合并范围和左上角值；未向其余格填造数据"
                        )
                    )
                units.append(u)
        finally:
            wb.close()
        extra, found = embedded_images(z, digest, source, objects, "xl")
        units.extend(extra)
        problems.extend(found)
        if any("/charts/" in p for p in z.namelist()):
            problems.append(
                issue("chart_not_rendered", "xl/charts", "图表未渲染，请结合底表核对", "blocking")
            )
        if any(p.startswith("xl/externalLinks/") for p in z.namelist()):
            problems.append(
                issue(
                    "external_workbook_not_fetched",
                    "xl/externalLinks",
                    "未获取外部工作簿，公式缓存需复核",
                    "blocking",
                )
            )
    return units, problems


def pdf(raw, digest, source, objects):
    units = []
    problems = []
    doc = pypdfium2.PdfDocument(raw)
    try:
        if len(doc) > 40:
            raise ValueError("PDF超过40页，请拆分")
        for index in range(len(doc)):
            page = doc[index]
            tp = page.get_textpage()
            try:
                text = tp.get_text_range()
                images = sum(1 for _ in page.get_objects(filter=[3]))
                path = objects / f"{digest}-page-{index + 1}.png"
                if not path.exists():
                    page.render(scale=2).to_pil().save(path)
                mode = (
                    "native"
                    if len(text.strip()) >= 20 and not images
                    else "mixed"
                    if text.strip()
                    else "scanned"
                )
                u = unit(
                    digest,
                    index + 1,
                    f"PDF第{index + 1}页 · {mode}",
                    source,
                    "pdf_page",
                    path=str(path),
                    pdf_mode=mode,
                    location=f"page:{index + 1}",
                    needs_ocr=mode != "native",
                    image_objects=images,
                )
                offset = 0
                for line in text.splitlines():
                    if line.strip():
                        u["native_lines"].append(
                            {
                                "text": line,
                                "raw_text": line,
                                "locator": f"page:{index + 1}/text-offset:{offset}",
                                "bbox": None,
                            }
                        )
                    offset += len(line) + 1
                if mode == "mixed":
                    problems.append(
                        issue("mixed_pdf_page", u["location"], "文字层与图像并存，不能因有文字而漏扫图像")
                    )
                if "\ufffd" in text:
                    problems.append(
                        issue(
                            "text_encoding_damage",
                            u["location"],
                            "原生文字层含替换字符，需对照渲染页",
                            "blocking",
                        )
                    )
                units.append(u)
            finally:
                tp.close()
                page.close()
    finally:
        doc.close()
    return units, problems


def ingest(filename, raw, objects):
    objects = Path(objects)
    objects.mkdir(parents=True, exist_ok=True)
    if len(raw) > 25 * 1024 * 1024:
        raise ValueError("单文件超过25MB")
    suffix = Path(filename).suffix.lower()
    digest = hashlib.sha256(raw).hexdigest()
    if suffix not in {".jpg", ".jpeg", ".png", ".pdf", ".docx", ".xlsx", ".doc", ".xls", ".docm", ".xlsm"}:
        raise ValueError("仅支持图片/PDF/DOCX/XLSX；其他类型未接收")
    source = objects / (digest + suffix)
    if not source.exists():
        source.write_bytes(raw)
    manifest = {
        "name": Path(filename).name,
        "sha256": digest,
        "bytes": len(raw),
        "format": suffix,
        "parser_version": VERSION,
        "source": str(source),
        "issues": [],
    }
    try:
        if suffix in {".doc", ".xls", ".docm", ".xlsm"}:
            raise ValueError("旧格式/宏文档仅存证，不执行；请转换为无宏DOCX/XLSX或PDF")
        if suffix in {".jpg", ".jpeg", ".png"}:
            units = [image_unit(raw, digest, 1, Path(filename).name, source, objects)]
            problems = []
        else:
            units, problems = {".docx": docx, ".xlsx": xlsx, ".pdf": pdf}[suffix](
                raw, digest, source, objects
            )
        if not units:
            raise ValueError("未找到可读取内容，不能标记为完成")
        manifest["technical_status"] = (
            "partial" if any(i["severity"] == "blocking" for i in problems) else "parsed"
        )
        manifest["issues"] = problems
    except Exception as e:
        message = type(e).__name__ + ": " + str(e)[:300]
        units = [unit(digest, 1, Path(filename).name, source, "unsupported")]
        manifest["technical_status"] = "failed"
        manifest["issues"] = [issue("intake_failed", filename, message, "blocking")]
    manifest.update(
        units=len(units),
        native_lines=sum(len(u["native_lines"]) for u in units),
        material_status="pending_human_completeness",
        business_status="not_confirmed",
    )
    for u in units:
        u["issues"] = [
            p
            for p in manifest["issues"]
            if p["location"] in {u.get("location"), u.get("part"), u.get("sheet")}
            or p["code"] == "intake_failed"
        ]
    return units, manifest
