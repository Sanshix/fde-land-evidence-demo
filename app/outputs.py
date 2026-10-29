"""Draft outputs from confirmed immutable snapshots. Never sends messages or pushes to other systems."""

import io
import json
import os
import re
import zipfile
from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

from . import business
from .domain import AS_OF, dec, money, to_upper_money

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
ET.register_namespace("w", W)
CODE_ROOT = Path(__file__).resolve().parents[1]
UNSET = "待主管部门配置"


def template(name):
    from .templates import CATALOG, active_path
    for kind, (filename, _) in CATALOG.items():
        if name == filename:
            return active_path(kind)
    return business.SAMPLES / "07_模板" / name


def notice_settings():
    """config/notice.json：缴费地点、核对地点、经办电话、期限天数、发文机关、文号前缀（赛题第八节：待主管部门确认）。"""
    path = Path(os.environ.get("FDE_NOTICE_CONFIG", CODE_ROOT / "config/notice.json"))
    data = json.loads(path.read_text()) if path.is_file() else {}
    return {
        "缴费地点": data.get("缴费地点") or UNSET,
        "核对地点": data.get("核对地点") or data.get("缴费地点") or UNSET,
        "经办电话": data.get("经办电话") or UNSET,
        "期限天数": int(data.get("期限天数") or 30),
        "发文机关": data.get("发文机关") or "××县自然资源局",
        "文号前缀": data.get("文号前缀") or "××自然资催",
        "configured": bool(data.get("缴费地点") and data.get("经办电话")),
    }


def paragraph_text(p):
    return "".join(t.text or "" for t in p.iter("{" + W + "}t"))


def replace_paragraph(p, values):
    texts = list(p.iter("{" + W + "}t"))
    original = "".join(t.text or "" for t in texts)
    updated = original
    for key, value in values.items():
        updated = updated.replace("【" + key + "】", str(value))
    if texts and updated != original:
        texts[0].text = updated
        for t in texts[1:]:
            t.text = ""


def set_paragraph(p, text):
    texts = list(p.iter("{" + W + "}t"))
    if texts:
        texts[0].text = text
        for t in texts[1:]:
            t.text = ""


def remove_element(root, element):
    for parent in root.iter():
        if element in list(parent):
            parent.remove(element)
            return


def year_ranges(years):
    """[2018,2019,2020,2023] → '2018至2020、2023'（短信模板允许写起止年度）。"""
    years = sorted(int(y) for y in years)
    out, start = [], None
    for i, y in enumerate(years):
        if start is None:
            start = y
        if i + 1 == len(years) or years[i + 1] != y + 1:
            out.append(str(start) if start == y else f"{start}至{y}")
            start = None
    return "、".join(out)


def arrears_rows(publication):
    return [r for r in publication["annual"] if dec(r["arrears"]) > 0]


def sms_text(publication, settings=None, deadline=None):
    """按 07_模板 短信模板生成；超过 70 字时把欠缴年度压缩为起止区间。"""
    settings = settings or notice_settings()
    years = [r["year"] for r in arrears_rows(publication)]
    when = f"请于{deadline}前" if deadline else f"请于收到通知后{settings['期限天数']}日内"
    org = settings["发文机关"]
    text = None
    for rendered in ("、".join(map(str, years)), year_ranges(years)):
        text = (
            f"【{org}】{publication['recipient']}您好，您使用的{publication['address']}国有土地尚欠{rendered}年度租金"
            f"共{publication['arrears']}元，{when}到{settings['缴费地点']}缴纳，咨询电话{settings['经办电话']}。"
        )
        if len(text) <= 70:
            break
    return {"text": text, "length": len(text), "segments": (len(text) + 66) // 67 if len(text) > 70 else 1}


def notice_docx(notice, publication, settings=None):
    settings = settings or notice_settings()
    p = publication
    fields = notice["contract_fields"]
    rows = arrears_rows(p)
    arrears = money(p["arrears"])  # 口径第八条：应缴之和 − 实缴之和（含多缴）
    overpaid = money(p.get("overpaid") or 0)
    values = {
        "年份": str(date.today().year),
        "编号": notice["id"][:6],
        "催缴对象": p["recipient"],
        "坐落": p["address"],
        "不动产单元号": p["registry_unit"] or "登记中无对应宗地",
        "面积": fields.get("area") or p.get("ledger", {}).get("登记面积", ""),
        "合同编号": fields.get("contract_no") or "待核实",
        "核算基准日": p["as_of"],
        "欠缴年度": year_ranges(r["year"] for r in rows),
        "欠缴金额": str(arrears),
        "欠缴金额大写": to_upper_money(arrears),
        "期限": settings["期限天数"],
        "缴费地点": settings["缴费地点"],
        "核对地点": settings["核对地点"],
        "经办电话": settings["经办电话"],
        "发文日期": "（审核签发后填写）",
        "到期日": fields.get("end", ""),
        "应缴合计": str(money(sum((dec(r["due"]) for r in rows), Decimal(0)))),
        "已缴合计": str(money(sum((dec(r["paid"]) for r in rows), Decimal(0)))),
        "欠缴合计": str(money(sum((dec(r["arrears"]) for r in rows), Decimal(0)))),
    }
    optional = {
        "续期": p.get("ledger", {}).get("合同状态") == "到期未续",
        "变更": any(i["type"] == "承租人已变更未办合同变更" for i in p.get("issues", [])),
    }
    output = io.BytesIO()
    with zipfile.ZipFile(template("国有土地租金催缴通知书（模板）.docx")) as src, zipfile.ZipFile(
        output, "w", zipfile.ZIP_DEFLATED
    ) as dest:
        for entry in src.infolist():
            raw = src.read(entry.filename)
            if entry.filename == "word/document.xml":
                root = ET.fromstring(raw)
                body = root.find("{" + W + "}body")
                # 明细表：只列欠缴年度（催缴通知面向产权人，已缴清年度不列入）
                for table in root.iter("{" + W + "}tbl"):
                    for row in list(table):
                        if "【年度】" in paragraph_text(row):
                            index = list(table).index(row)
                            table.remove(row)
                            for offset, item in enumerate(rows):
                                copy = deepcopy(row)
                                for paragraph in copy.iter("{" + W + "}p"):
                                    replace_paragraph(
                                        paragraph,
                                        {"年度": item["year"], "应缴": item["due"], "已缴": item["paid"], "欠缴": item["arrears"]},
                                    )
                                table.insert(index + offset, copy)
                for paragraph in list(root.iter("{" + W + "}p")):
                    text = paragraph_text(paragraph)
                    if text.startswith("【可选段落："):
                        keep = ("续期" in text and optional["续期"]) or ("变更" in text and optional["变更"])
                        if keep:
                            set_paragraph(paragraph, re.sub(r"^【可选段落：|】$", "", text))
                        else:
                            remove_element(root, paragraph)
                        continue
                    if text.startswith("短信通知模板") or text.startswith("【××县自然资源局】") or text.startswith("短信不超过"):
                        remove_element(root, paragraph)  # 短信模板单独生成，不进通知书
                        continue
                    replace_paragraph(paragraph, values)
                if overpaid > 0:
                    label = ET.Element("{" + W + "}p")
                    run = ET.SubElement(label, "{" + W + "}r")
                    t = ET.SubElement(run, "{" + W + "}t")
                    t.text = f"注：另有多缴{overpaid}元，已在欠缴金额中抵减（附表按年度列示，合计与欠缴金额的差额即为多缴抵减）。"
                    section = body.find("{" + W + "}sectPr")
                    if section is None:
                        body.append(label)
                    else:
                        body.insert(list(body).index(section), label)
                if not settings["configured"]:
                    label = ET.Element("{" + W + "}p")
                    run = ET.SubElement(label, "{" + W + "}r")
                    t = ET.SubElement(run, "{" + W + "}t")
                    t.text = "草稿：缴费地点/经办电话尚未在 config/notice.json 配置，不得发出"
                    body.insert(0, label)
                raw = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            dest.writestr(entry, raw)
    return output.getvalue()


def quarter_bounds(year, quarter):
    start = date(year, (quarter - 1) * 3 + 1, 1)
    end = date(year, quarter * 3 + 1, 1) if quarter < 4 else date(year + 1, 1, 1)
    return start, end


def tax_workbook(ledgers, docs, source, year, quarter):
    """税务推送单。“本期应缴”= 缴纳期限（6月30日）落在本期的年度应缴，与模板示例（第二季度）一致；口径待税务部门确认。"""
    if not 1998 <= year <= AS_OF.year or quarter not in (1, 2, 3, 4):
        raise ValueError("推送期间无效")
    start, end = quarter_bounds(year, quarter)
    if start > AS_OF:
        raise ValueError("推送期间晚于核算基准日")
    wb = openpyxl.load_workbook(template("税务推送单（字段与样式）.xlsx"))
    ws = wb["推送单"]
    ws.delete_rows(2, ws.max_row)
    by_id = {d["id"]: d for d in docs}
    count = 0
    for pub in ledgers:
        p = pub["payload"]
        document = by_id.get(p.get("latest_document"))
        if not p["documents"] or not document:
            continue
        fields = document["review"]["fields"]
        chosen = set()
        for identifier in p["documents"]:
            if identifier in by_id:
                chosen.update(str(x) for x in by_id[identifier]["review"].get("payment_ids", []))
        paid = {}
        for r in source["payments"]:
            day = date.fromisoformat(str(r["缴费日期"])[:10])
            if str(r["序号"]) in chosen and start <= day < end:
                paid[str(r["票据号"])] = r
        for event in source.get("business_events", []):
            if event["kind"] == "payment" and set(event["documents"]) & set(p["documents"]):
                payment = event["payment"]
                if start <= date.fromisoformat(payment["缴费日期"]) < end:
                    paid[str(payment["票据号"])] = payment
        current = next((r for r in source["registry"] if r["不动产单元号"] == p["registry_unit"]), None)
        annual = next((r["due"] for r in p["annual"] if r["year"] == year), "0.00")
        deadline_in_period = start <= date(year, 6, 30) < end
        count += 1
        ws.append(
            [
                count,
                p["recipient"],
                str(current["证件号码"]) if current else fields.get("identity", ""),
                p["registry_unit"],
                p["address"],
                float(dec(fields.get("area") or 0)),
                (fields.get("use") or "") + "用地" if fields.get("use") else "",
                fields.get("start", ""),
                fields.get("end", ""),
                float(dec(p["segments"][-1]["annual"])) if p.get("segments") else 0.0,
                float(dec(annual)) if deadline_in_period else 0.0,
                float(money(sum((dec(r["金额（元）"]) for r in paid.values()), Decimal(0)))),
                "；".join(sorted({str(r["缴费日期"])[:10] for r in paid.values()})),
                "；".join(sorted(paid)),
                f"{year}年第{quarter}季度",
                f"本期应缴按6月30日缴纳期限归入第二季度；台账版本{pub['id']}",
            ]
        )
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                cell.value = "'" + cell.value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), count


def contract_docx(values, kind, original_no=None):
    """按制式模板生成新签/续期合同草稿（表2第6项）。values 的键为模板【】内名称。"""
    output = io.BytesIO()
    with zipfile.ZipFile(template("国有建设用地使用权租赁合同（模板）.docx")) as src, zipfile.ZipFile(
        output, "w", zipfile.ZIP_DEFLATED
    ) as dest:
        for entry in src.infolist():
            raw = src.read(entry.filename)
            if entry.filename == "word/document.xml":
                root = ET.fromstring(raw)
                for paragraph in list(root.iter("{" + W + "}p")):
                    text = paragraph_text(paragraph)
                    if text.startswith("说明："):
                        if kind == "续期" and original_no:
                            set_paragraph(paragraph, f"本合同为原合同（编号：{original_no}）期满续签。")
                        else:
                            remove_element(root, paragraph)
                        continue
                    if text == "【固定 / 随基准地价调整】":
                        set_paragraph(paragraph, values.get("计租方式", "固定"))
                        continue
                    replace_paragraph(paragraph, values)
                raw = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            dest.writestr(entry, raw)
    return output.getvalue()
