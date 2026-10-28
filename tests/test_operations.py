import hashlib
import json
import sqlite3
from pathlib import Path

import pytest
from test_api import HEAD, new_doc, preview, reviewed
from test_api import client as client_fixture

from app import store
from app.maintenance import restore_backup

client = client_fixture


def test_health_and_restore_checks_contents(client, tmp_path):
    doc, _ = new_doc(client)
    body = preview(client, doc, reviewed(doc))
    before = client.post(f"/api/documents/{doc['id']}/confirm", json=body, headers=HEAD).json()
    backup = client.post("/api/backup", json={}, headers=HEAD | {"X-FDE-Role": "admin"}).json()
    target = tmp_path / "restore-copy"
    restored = restore_backup(backup["path"], target)
    assert restored["counts"]["documents"] == 1 and restored["counts"]["publications"] == 1
    db = sqlite3.connect(target / "fde.sqlite")
    assert db.execute("SELECT digest FROM publications").fetchone()[0] == before["digest"]
    pages = json.loads(db.execute("SELECT pages FROM documents").fetchone()[0])
    assert Path(pages[0]["path"]).is_relative_to(target) and Path(pages[0]["path"]).exists()
    intake = json.loads(db.execute("SELECT intake FROM documents").fetchone()[0])
    assert Path(intake["files"][0]["source"]).is_relative_to(target)
    assert Path(intake["files"][0]["source"]).exists()
    db.close()
    assert client.get("/healthz").json()["schema_version"] == 6
    with pytest.raises(ValueError):
        restore_backup(backup["path"], target)
    (Path(backup["path"]) / "objects" / Path(pages[0]["path"]).name).write_bytes(b"corrupted")
    with pytest.raises(ValueError):
        restore_backup(backup["path"], tmp_path / "corrupt")


def test_rule_change_transaction_and_historical_notice(client):
    doc, _ = new_doc(client)
    b = reviewed(doc)
    b["fields"].update(mode="随标准调整", grade="一级", annual="0")
    b = preview(client, doc, b)
    before = client.post(f"/api/documents/{doc['id']}/confirm", json=b, headers=HEAD).json()
    assert "id" in before, before
    notice = client.post(f"/api/notices/{before['id']}", json={}, headers=HEAD).json()
    config = client.get("/api/rules").json()
    new = config["rates"]
    for row in new:
        row["一级"] = str(float(row["一级"]) + 1)
    admin = HEAD | {"X-FDE-Role": "admin"}
    assert client.post("/api/rules/preview", json={"rates": new}, headers=HEAD).status_code == 403
    draft = client.post("/api/rules/preview", json={"rates": new}, headers=admin).json()
    assert draft["impact"][0]["before"]["payload"]["due"] != draft["impact"][0]["after"]["due"]
    assert (
        client.post(
            "/api/rules/confirm",
            json={"rates": new, "reason": "test", "draft_hash": "wrong", "confirm_calculation": True},
            headers=admin,
        ).status_code
        == 409
    )
    assert (
        client.post(
            "/api/rules/confirm",
            json={"rates": new, "reason": "test", "draft_hash": draft["draft_hash"]},
            headers=admin,
        ).status_code
        == 400
    )
    assert (
        client.post(
            "/api/rules/confirm",
            json={
                "rates": new,
                "reason": "synthetic rule test",
                "draft_hash": draft["draft_hash"],
                "confirm_calculation": True,
            },
            headers=admin,
        ).status_code
        == 200
    )
    after = client.get("/api/ledger").json()[0]
    assert after["payload"]["rate_version"] != before["payload"]["rate_version"]
    with store.connection() as db:
        assert (
            json.loads(db.execute("SELECT payload FROM notices WHERE id=?", (notice["id"],)).fetchone()[0])[
                "arrears"
            ]
            == before["payload"]["arrears"]
        )
        assert db.execute("SELECT COUNT(*) FROM publications").fetchone()[0] == 2


def test_actual_accounts_ignore_forged_role_headers(client, tmp_path, monkeypatch):
    salt = "1234567890abcdef1234567890abcdef"
    password = "synthetic-test-password"
    hashvalue = hashlib.scrypt(password.encode(), salt=bytes.fromhex(salt), n=16384, r=8, p=1).hex()
    path = tmp_path / "users.json"
    path.write_text(json.dumps({"reader": {"salt": salt, "hash": hashvalue, "role": "viewer"}}))
    monkeypatch.setenv("FDE_USERS_FILE", str(path))
    assert client.get("/api/status").status_code == 401
    assert client.get("/api/status", auth=("reader", password)).status_code == 200
    assert (
        client.post(
            "/api/backup", json={}, headers=HEAD | {"X-FDE-Role": "admin"}, auth=("reader", password)
        ).status_code
        == 403
    )


def test_batch_and_output_drafts(client):
    import io
    import zipfile

    from PIL import Image

    raw = io.BytesIO()
    Image.new("RGB", (120, 100), "blue").save(raw, format="PNG")
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("TEST-A_1.png", raw.getvalue())
    response = client.post(
        "/api/import-batch",
        files={"file": ("batch.zip", archive.getvalue(), "application/zip")},
        headers=HEAD,
    )
    assert response.status_code == 200, response.text
    identifier = response.json()["documents"][0]["id"]
    doc = client.get("/api/documents/" + identifier).json()
    body = preview(client, doc, reviewed(doc))
    pub = client.post(f"/api/documents/{identifier}/confirm", json=body, headers=HEAD).json()
    n = client.post(f"/api/notices/{pub['id']}", json={}, headers=HEAD).json()
    output = client.get(f"/api/notices/{n['id']}/document")
    with zipfile.ZipFile(io.BytesIO(output.content)) as z:
        text = z.read("word/document.xml").decode()
        assert "草稿" in text and "120.00" in text and "【" not in text
    push = client.post("/api/tax-push", json={"year": 2026, "quarter": 3}, headers=HEAD)
    assert push.status_code == 200, push.text
    output = client.get(f"/api/tax-push/{push.json()['id']}/document")
    assert output.status_code == 200
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(output.content))
    # 第三季度不含 6 月 30 日缴纳期限，本期应缴为 0；每次推送留存可追溯
    assert wb["推送单"].max_row == 2 and wb["推送单"].cell(2, 11).value == 0
    assert client.get("/api/tax-pushes").json()[0]["rows"] == 1
    batch = client.get("/api/status").json()["batches"][0]
    assert batch["all_published"] and batch["within_20_minutes"]
    with store.connection() as db:
        with pytest.raises(sqlite3.IntegrityError):
            db.execute("UPDATE publications SET digest=?", ("forged",))


def test_rule_change_invalidates_unpublished_calculation(client):
    doc, _ = new_doc(client)
    body = preview(client, doc, reviewed(doc))
    config = client.get("/api/rules").json()
    rates = config["rates"]
    rates[0]["一级"] = str(float(rates[0]["一级"]) + 0.01)
    admin = HEAD | {"X-FDE-Role": "admin"}
    impact = client.post("/api/rules/preview", json={"rates": rates}, headers=admin).json()
    changed = client.post(
        "/api/rules/confirm",
        json={
            "rates": rates,
            "reason": "synthetic changed-source test",
            "draft_hash": impact["draft_hash"],
            "confirm_calculation": True,
        },
        headers=admin,
    )
    assert changed.status_code == 200
    assert client.post(f"/api/documents/{doc['id']}/confirm", json=body, headers=HEAD).status_code == 409
    assert client.get("/api/ledger").json() == []


def test_explicit_recovery_retry_keeps_original_batch_clock(client):
    import time

    doc, _ = new_doc(client)
    old = time.time() - 1800
    with store.connection() as db:
        db.execute("UPDATE jobs SET state='failed'")
        db.execute("UPDATE batches SET received=?", (old,))
    started = time.time()
    response = client.post(
        f"/api/documents/{doc['id']}/retry", json={"version": doc["version"]}, headers=HEAD
    )
    assert response.status_code == 200
    with store.connection() as db:
        job = db.execute("SELECT created FROM jobs WHERE id=?", (response.json()["job"],)).fetchone()[0]
        assert job >= started
        assert db.execute("SELECT received FROM batches").fetchone()[0] == old
    batch = client.get("/api/status").json()["batches"][0]
    assert batch["elapsed_seconds"] >= 1800 and not batch["within_20_minutes"]


def test_restore_native_units_without_rendered_image(client, tmp_path, monkeypatch):
    from test_governance import make_docx

    response = client.post(
        "/api/import",
        files={"files": ("native.docx", make_docx(["土地面积：10"]), "application/octet-stream")},
        data={"name": "native-only-restore"},
        headers=HEAD,
    )
    assert response.status_code == 200
    backup = client.post("/api/backup", json={}, headers=HEAD | {"X-FDE-Role": "admin"}).json()
    target = tmp_path / "native-restored"
    result = restore_backup(backup["path"], target)
    assert result["integrity"] == "ok" and not (target / "RESTORE_INCOMPLETE").exists()
    db = sqlite3.connect(target / "fde.sqlite")
    page = json.loads(db.execute("SELECT pages FROM documents").fetchone()[0])[0]
    assert page["path"] is None and Path(page["source"]).is_file()
    db.close()
    (target / "RESTORE_INCOMPLETE").write_text("synthetic failed restoration marker")
    monkeypatch.setattr(store, "DATA", target)
    with pytest.raises(RuntimeError, match="恢复未完成"):
        store.init()


def test_restore_does_not_require_original_alias_to_exist(client, tmp_path):
    from app.maintenance import create_backup

    doc, _ = new_doc(client)
    alias = tmp_path.parent / (tmp_path.name + "-alias")
    try:
        alias.symlink_to(tmp_path, target_is_directory=True)
    except OSError:
        pytest.skip("OS does not allow this symlink fixture")
    with store.connection() as db:
        for column in ("pages", "intake"):
            text = db.execute(f"SELECT {column} FROM documents").fetchone()[0]
            db.execute(f"UPDATE documents SET {column}=?", (text.replace(str(tmp_path), str(alias)),))
    backup = tmp_path / "alias-backup"
    create_backup(alias, backup, store.DB)
    alias.unlink()
    target = tmp_path / "restored-without-alias"
    assert restore_backup(backup, target)["integrity"] == "ok"
    restored = sqlite3.connect(target / "fde.sqlite")
    page = json.loads(restored.execute("SELECT pages FROM documents").fetchone()[0])[0]
    assert Path(page["source"]).is_file() and Path(page["source"]).is_relative_to(target)
    restored.close()
