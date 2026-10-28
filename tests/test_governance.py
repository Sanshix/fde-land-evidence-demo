import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from PIL import Image
from test_api import HEAD, new_doc, preview, reviewed
from test_api import client as client_fixture

from app import governance, intake
from app.cloud_cascade import run_cloud_cascade

client = client_fixture


def make_docx(texts, image=False):
    stream = io.BytesIO()
    root = ET.Element("{" + intake.W + "}document")
    body = ET.SubElement(root, "{" + intake.W + "}body")
    for text in texts:
        p = ET.SubElement(body, "{" + intake.W + "}p")
        r = ET.SubElement(p, "{" + intake.W + "}r")
        t = ET.SubElement(r, "{" + intake.W + "}t")
        t.text = text
    with zipfile.ZipFile(stream, "w") as z:
        z.writestr("word/document.xml", ET.tostring(root))
        if image:
            picture = io.BytesIO()
            Image.new("RGB", (100, 100), "white").save(picture, format="PNG")
            z.writestr("word/media/image1.png", picture.getvalue())
    return stream.getvalue()


def test_native_docx_evidence_and_embedded_image(tmp_path):
    raw = make_docx(
        [
            "合同编号：TEST-NATIVE",
            "承租人：模拟人",
            "土地面积：１２３.４５平方米",
            "租赁起始日：2026年1月1日",
            "计租方式：固定  随标准调整",
        ],
        image=True,
    )
    units, manifest = intake.ingest("sample.docx", raw, tmp_path)
    assert manifest["technical_status"] == "parsed"
    assert len(units) == 2 and units[1]["kind"] == "image"
    assert Path(manifest["source"]).read_bytes() == raw
    cs = governance.native_candidates(units)
    assert cs["area"][0]["value"] == "123.45" and cs["area"][0]["raw_value"] == "１２３.４５平方米"
    assert cs["start"][0]["value"] == "2026-01-01"
    assert not cs["mode"]
    result = run_cloud_cascade(units, tmp_path / "run", config={"enabled": False})
    assert result["mode"] == "unconfigured" and result["external_calls"] == 0
    assert result["candidates"]["area"][0]["source_locator"].endswith("/p3")


def test_normalization_never_drops_negative_sign(tmp_path):
    units, _ = intake.ingest(
        "negative.docx", make_docx(["年租金：－１２０元", "土地面积：−10平方米"]), tmp_path
    )
    candidates = governance.native_candidates(units)
    assert candidates["annual"][0]["value"] == "-120"
    assert candidates["annual"][0]["raw_value"] == "－１２０元"
    assert candidates["area"][0]["value"] == "-10"


def test_excel_precision_formula_hidden_and_merges(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "合同"
    ws.append(["宗地号", "00123"])
    ws["B1"].number_format = "@"
    ws.append(["面积", 12])
    ws["B2"].number_format = "00000"
    ws["B3"] = "=B2*10"
    ws["B4"] = 123456789012345678
    ws.merge_cells("A5:B5")
    ws["A5"] = "合并标题"
    hidden = wb.create_sheet("隐藏证据")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "隐藏内容"
    ws.row_dimensions[6].hidden = True
    ws["A6"] = "隐藏单元格"
    stream = io.BytesIO()
    wb.save(stream)
    units, manifest = intake.ingest("test.xlsx", stream.getvalue(), tmp_path)
    cells = {r["cell"]: r for r in units[0]["native_lines"]}
    assert cells["B1"]["text"] == "00123"
    assert cells["B2"]["raw_xml_value"] == "12" and cells["B2"]["display_value"] == "00012"
    assert cells["B3"]["formula"] == "B2*10" and cells["B3"]["cached_value"] is None
    assert cells["B4"]["raw_xml_value"] is not None
    assert units[1]["hidden"] and units[0]["merged_ranges"] == ["A5:B5"]
    codes = {p["code"] for p in manifest["issues"]}
    assert {
        "hidden_sheet",
        "hidden_cell",
        "formula_not_executed",
        "numeric_identifier_precision",
        "merged_cells",
    } <= codes
    assert manifest["technical_status"] == "partial"


def test_failed_formats_remain_visible(client):
    response = client.post(
        "/api/import",
        files={"files": ("old.doc", b"NOT AN EXECUTABLE MACRO", "application/msword")},
        data={"name": "unsupported-old-doc"},
        headers=HEAD,
    )
    assert response.status_code == 200
    doc = client.get("/api/documents/" + response.json()["id"]).json()
    assert doc["intake"]["technical_status"] == "failed"
    assert any(p["code"] == "intake_failed" for p in doc["governance_issues"])
    queue = client.get("/api/governance").json()
    assert queue[0]["id"] == doc["id"]
    body = preview(client, doc, reviewed(doc))
    assert client.post(f"/api/documents/{doc['id']}/confirm", json=body, headers=HEAD).status_code == 400


def test_draft_cross_material_evidence_and_stale_confirmation(client):
    doc, _ = new_doc(client)
    body = preview(client, doc, reviewed(doc))
    assert (
        client.post(f"/api/documents/{doc['id']}/draft", json=body, headers=HEAD).json()["effective"] is False
    )
    assert client.get(f"/api/documents/{doc['id']}").json()["draft"]["fields"]["area"] == "10"
    attached = client.post(
        f"/api/documents/{doc['id']}/attach",
        files={
            "files": (
                "supplement.docx",
                make_docx(["土地面积：10"]),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"version": doc["version"]},
        headers=HEAD,
    )
    assert attached.status_code == 200, attached.text
    assert attached.json()["confirmation"] == "invalidated"
    assert client.post(f"/api/documents/{doc['id']}/confirm", json=body, headers=HEAD).status_code == 409
    current = client.get(f"/api/documents/{doc['id']}").json()
    assert current["draft"] is None and len(current["pages"]) == 2
    updated = reviewed(current)
    native = current["pages"][1]
    updated["citations"]["area"] = native["id"]
    updated = preview(client, current, updated)
    assert client.post(f"/api/documents/{doc['id']}/confirm", json=updated, headers=HEAD).status_code == 400
    updated["citation_details"] = {
        "area": [
            {"unit_id": native["id"], "locator": native["native_lines"][0]["locator"]},
            {"unit_id": current["pages"][0]["id"], "locator": ""},
        ]
    }
    updated = preview(client, current, updated)
    assert client.post(f"/api/documents/{doc['id']}/confirm", json=updated, headers=HEAD).status_code == 200


def test_preview_binds_nonfinancial_fields_and_sources(client):
    doc, _ = new_doc(client)
    body = preview(client, doc, reviewed(doc))
    changed = body | {"fields": body["fields"] | {"identity": "new-identity"}}
    changed["citations"]["identity"] = doc["pages"][0]["id"]
    assert client.post(f"/api/documents/{doc['id']}/confirm", json=changed, headers=HEAD).status_code == 409


def test_actual_existing_word_and_excel_read_only(tmp_path):
    from app.business import SAMPLES as root

    files = [root / "06_计租标准与核算口径.docx", root / "02_不动产登记导出.xlsx"]
    for f in files:
        before = f.read_bytes()
        units, manifest = intake.ingest(f.name, before, tmp_path)
        assert manifest["technical_status"] in {"parsed", "partial"}
        assert manifest["native_lines"] > 10 and units
        assert f.read_bytes() == before


def pdf_fixture(mode):
    import zlib

    content = b"BT /F1 14 Tf 20 150 Td (Synthetic native contract 12345) Tj ET" if mode != "scanned" else b""
    if mode != "native":
        content += b" q 80 0 0 80 20 30 cm /Im0 Do Q"
    resources = b"/Font << /F1 4 0 R >>" + (b" /XObject << /Im0 6 0 R >>" if mode != "native" else b"")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Resources << "
        + resources
        + b" >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    if mode != "native":
        pixels = zlib.compress(bytes([230, 230, 230]) * 16)
        objs.append(
            b"<< /Type /XObject /Subtype /Image /Width 4 /Height 4 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode /Length "
            + str(len(pixels)).encode()
            + b" >>\nstream\n"
            + pixels
            + b"\nendstream"
        )
    result = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objs, 1):
        offsets.append(len(result))
        result += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    start = len(result)
    result += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode() + b"".join(
        f"{n:010} 00000 n \n".encode() for n in offsets[1:]
    )
    return result + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode()


def test_pdf_native_scanned_and_mixed_are_distinct(tmp_path):
    for mode in ("native", "scanned", "mixed"):
        units, manifest = intake.ingest(mode + ".pdf", pdf_fixture(mode), tmp_path)
        assert manifest["technical_status"] == "parsed", manifest
        assert units[0]["pdf_mode"] == mode
        assert units[0]["needs_ocr"] == (mode != "native")
        assert bool(units[0]["native_lines"]) == (mode != "scanned")
        assert Path(units[0]["path"]).is_file()
    result = run_cloud_cascade(
        intake.ingest("native.pdf", pdf_fixture("native"), tmp_path)[0],
        tmp_path / "run-native",
        config={"enabled": False},
    )
    assert result["mode"] == "native_only" and result["external_calls"] == 0


def test_macro_and_xml_entity_never_execute(tmp_path):
    units, m = intake.ingest("macro.xlsm", b"not executed", tmp_path)
    assert m["technical_status"] == "failed" and units[0]["kind"] == "unsupported"
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as z:
        z.writestr("word/document.xml", b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>')
    _, m = intake.ingest("entity.docx", stream.getvalue(), tmp_path)
    assert m["technical_status"] == "failed"
    assert "DTD" in m["issues"][0]["message"]
